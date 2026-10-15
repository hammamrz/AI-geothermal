#!/usr/bin/env python3
"""Membangun PPTX technical review geothermal bergaya PLN IP dari JSON terstruktur.

Struktur deck mengikuti tulang punggung Technical Review / Comment Sheet Report:
cover -> ringkasan eksekutif -> agenda -> pengendalian dokumen (opsional) -> tujuan/lingkup ->
dokumen yang dikaji -> metodologi -> sebaran temuan -> daftar temuan -> rincian temuan ->
risiko antardisiplin -> data gap -> langkah penyelesaian -> kesimpulan dan langkah selanjutnya ->
sumber -> penutup.

Setiap item review disajikan berurutan: DATA EKSISTING (apa yang tertulis/tersedia di dokumen,
apa adanya) -> KOMENTAR REVIEW (penilaian reviewer + dasar KB) -> LANGKAH PENYELESAIAN (langkah
konkret bernomor: tindakan, keluaran yang bisa diperiksa, pihak yang mengerjakan). Deck tidak
menampilkan persentase/status penyelesaian kajian.

Tampilan mengikuti framework skill presentasi-pln-ip dan presentasi-tvv (pustaka di
assets/pustaka): judul dua warna (topik biru + sub-topik hitam) tanpa label kecil di atasnya,
tanpa garis aksen, tanpa sudut membulat/bayangan/gradien, tanpa deretan kartu; tabel
bersumbu; model grafis asli PowerPoint (deret angka kunci, alur proses, chart bertumpuk,
matriks risiko 5x5). Teks mendukung *miring* dan **tebal**.

Pemakaian:
    python build_review_ppt.py review.json output.pptx

Format JSON (kunci lama tetap diterima):
{
  "report_title": "Kajian Teknis Drilling Program Sumur X",
  "project_name": "Lapangan / proyek",
  "report_no": "TR-001", "revision": "0", "date": "30 September 2026",
  "prepared_by": "...", "reviewed_by": "...", "approved_by": "...",
  "executive_summary": ["...", "..."],
  "objective": "...", "scope": ["..."], "review_basis": ["..."], "limitations": ["..."],
  "methodology": ["..."],                              # opsional, 3-5 langkah
  "documents_reviewed": [{"title","number","revision","date","discipline","scope"}],
  "findings": [{"id","discipline","location","title",
                "existing": "data eksisting apa adanya dari dokumen (angka, asumsi, metode)",
                "finding": "komentar review: apa yang kurang/tidak konsisten/berisiko",
                "comment": "rekomendasi singkat", "impact","basis","source","priority","confidence",
                "resolution_steps": ["Hitung ulang ...", {"step": "...", "pic": "...",
                                     "due": "...", "evidence": "..."}],
                "image_path"}],
  "key_risks": [{"interface","risk","impact","mitigation","related_findings",
                 "likelihood": 1-5, "severity": 1-5}],   # likelihood+severity -> matriks 5x5
  "data_gaps": [{"id","gap","impact","required","owner","resolution_steps": ["..."]}],
  "actions": [{"id","finding_id","action","pic","due","evidence"}],  # dipakai bila temuan tanpa
                                                                    # resolution_steps
  "conclusion": "...", "next_steps": ["langkah prioritas 1", "..."],  # alias: conditions_to_proceed
  "overall_status": "Perlu revisi (ditampilkan sebagai rekomendasi)", "references_used": ["..."]
}
"""
from __future__ import annotations

import json
import os
import re
import sys
from pathlib import Path
from typing import Any

PUSTAKA = Path(__file__).resolve().parent.parent / "assets" / "pustaka"
sys.path.insert(0, str(PUSTAKA))

import plnip_deck as P  # noqa: E402  palet, teks, bullet, judul kolom
import plnip_grafis as G  # noqa: E402  model grafis
import tvv_deck as T  # noqa: E402  tabel berpenanda status, matriks risiko, gambar

from pptx import Presentation  # noqa: E402
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR  # noqa: E402
from pptx.util import Inches  # noqa: E402

SW, SH = 13.333, 7.5
ASSET_DIR = Path(__file__).resolve().parent.parent / "references" / "master_assets"
MASTER_COVER = ASSET_DIR / "MASTER_COVER.jpg"
MASTER_CONTENT = ASSET_DIR / "MASTER_CONTENT.jpg"
MASTER_CLOSING = ASSET_DIR / "MASTER_CLOSING.jpg"

# Grid untuk master isi geothermal (logo Danantara + PLN IP di kanan atas mulai x ±9,3)
L, R = 0.55, 12.78
CW = R - L
TITLE_X, TITLE_Y, TITLE_W = 0.55, 0.32, 8.45
BODY_TOP, BODY_BOTTOM = 1.4, 6.75
LEFT_X, LEFT_W = L, 6.55
RIGHT_X, RIGHT_W = 7.5, R - 7.5
TABLE_PT = 11
_PAGED: set[int] = set()  # slide_id slide isi yang diberi nomor halaman


# ----------------------------------------------------------------- data
def as_list(value: Any) -> list[Any]:
    if value is None:
        return []
    if isinstance(value, (list, tuple)):
        return list(value)
    return [value]


def text_of(value: Any, sep: str = "; ") -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, (list, tuple)):
        return sep.join(str(x) for x in value if x not in (None, ""))
    return str(value)


def compact(value: Any, limit: int = 180) -> str:
    s = " ".join(text_of(value).split())
    if len(s) <= limit:
        return s
    s = s[: max(0, limit - 1)].rstrip()
    if " " in s[limit // 2:]:  # potong di batas kata, bukan di tengah kata
        s = s[: s.rfind(" ")].rstrip(" ,;:")
    if s.count("*") % 2:  # jangan memotong markup miring di tengah
        s = s[: s.rfind("*")].rstrip()
    return s + "…"


def report_title(d: dict[str, Any]) -> str:
    return d.get("report_title") or d.get("title") or "Kajian Teknis Geothermal"


def project_name(d: dict[str, Any]) -> str:
    return d.get("project_name") or d.get("subtitle") or ""


def executive_messages(d: dict[str, Any]) -> list[str]:
    msgs = d.get("executive_summary") or d.get("summary_messages") or []
    if isinstance(msgs, str):
        parts = [x.strip() for x in msgs.replace("\r", "\n").split("\n") if x.strip()]
        return parts or [msgs]
    return [text_of(x) for x in as_list(msgs)]


def normalize_finding(f: dict[str, Any]) -> dict[str, Any]:
    return {
        **f,
        "discipline": f.get("discipline", ""),
        "location": f.get("location") or f.get("target_location") or "",
        "existing": f.get("existing") or f.get("existing_data") or f.get("data_eksisting") or "",
        "comment": f.get("comment") or f.get("recommendation") or "",
        "basis": f.get("basis", ""),
        "source": f.get("source") or f.get("technical_basis") or "",
        "impact": f.get("impact") or f.get("potential_impact") or "",
    }


def tanpa_tindakan(f: dict[str, Any]) -> bool:
    """Item yang menurut reviewer sudah memadai (komentar 'tidak ada tindakan')."""
    return bool(re.search(r"(?i)tidak (ada|perlu) tindak|no (further )?action", text_of(f.get("comment"))))


def steps_of(f: dict[str, Any], d: dict[str, Any] | None = None) -> list[dict[str, str]]:
    """Langkah penyelesaian suatu temuan: resolution_steps; bila kosong, actions yang merujuk ID
    temuan; bila masih kosong, rekomendasi (comment) sebagai satu langkah."""
    raw = as_list(f.get("resolution_steps") or f.get("langkah_penyelesaian") or f.get("steps"))
    steps = []
    for r in raw:
        if isinstance(r, dict):
            steps.append({"step": text_of(r.get("step") or r.get("action") or r.get("langkah")),
                          "pic": text_of(r.get("pic") or r.get("owner")), "due": text_of(r.get("due")),
                          "evidence": text_of(r.get("evidence") or r.get("closure_evidence"))})
        elif text_of(r):
            steps.append({"step": text_of(r), "pic": "", "due": "", "evidence": ""})
    if not steps and d is not None and f.get("id"):
        for a in as_list(d.get("actions")):
            if isinstance(a, dict) and f["id"] in text_of(a.get("finding_id") or a.get("related_findings")):
                steps.append({"step": text_of(a.get("action")), "pic": text_of(a.get("pic") or a.get("owner")),
                              "due": text_of(a.get("due") or a.get("timing")),
                              "evidence": text_of(a.get("evidence") or a.get("closure_evidence"))})
    if not steps and f.get("comment") and not tanpa_tindakan(f):
        steps.append({"step": text_of(f["comment"]), "pic": "", "due": "", "evidence": ""})
    return [x for x in steps if x["step"]]


def numbered(items: list[str], limit: int = 150) -> list[str]:
    return [f"{i}. {compact(t, limit)}" for i, t in enumerate(items, 1)]


def findings_of(d: dict[str, Any]) -> list[dict[str, Any]]:
    return [normalize_finding(f) for f in as_list(d.get("findings")) if isinstance(f, dict)]


def priority_level(priority: str) -> str:
    """Kenali prioritas dalam bahasa Inggris maupun Indonesia."""
    p = (priority or "").strip().lower()
    if p.startswith(("h", "critical", "major", "tinggi", "kritis")):
        return "high"
    if p.startswith(("m", "sedang")):
        return "medium"
    return "low"


# ----------------------------------------------------------------- slide dasar
def background(slide, path: Path):
    if not path.exists():
        raise FileNotFoundError(f"Aset master tidak ditemukan: {path}")
    slide.shapes.add_picture(str(path), 0, 0, width=Inches(SW), height=Inches(SH))


def content_slide(prs, topik: str, subtopik: str = ""):
    """Slide isi: master geothermal + judul dua warna (topik biru, sub-topik hitam)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide, MASTER_CONTENT)
    teks = f"{topik} {subtopik}".strip()
    size = 22
    if P._title_lines(teks, TITLE_W, size) > 2:
        size = 18
    box = slide.shapes.add_textbox(Inches(TITLE_X), Inches(TITLE_Y), Inches(TITLE_W), Inches(0.9))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    P._fill_runs(p, topik, size, P.PRIMARY, bold=True)
    if subtopik:
        P._fill_runs(p, " " + subtopik, size, P.INK, bold=True)
    _PAGED.add(slide.slide_id)
    return slide


def add_source(slide, teks: str):
    """Satu baris sumber (pola TVV) untuk slide yang memuat bukti teknis."""
    if teks:
        T.add_sumber(slide, compact(teks, 200))


def caption(slide, teks: str, x: float = L, w: float = CW):
    """Keterangan singkat cara membaca grafis, di bawah konten (pola PLN IP)."""
    P.add_keterangan_bawah(slide, teks, x=x, w=w, size=12)


def clean(teks: Any) -> str:
    """Simbol panah/dekoratif dari input diganti tanda baca biasa (aturan anti-dekorasi)."""
    out = text_of(teks)
    for sym, rep_ in (("↔", " / "), ("→", " ke "), ("⇒", " ke "), ("✅", ""), ("⏳", ""), ("❌", "")):
        out = out.replace(sym, rep_)
    return " ".join(out.split())


def table_pages(rows: list, per_page: int) -> list[list]:
    return [rows[i:i + per_page] for i in range(0, len(rows), per_page)] or [[]]


def text_height(teks: str, w_in: float, pt: float) -> float:
    """Perkiraan tinggi teks (inci) agar bagian berikutnya tidak berjarak terlalu jauh."""
    chars = max(20, int(w_in * 72 / (pt * 0.5)))
    lines = max(1, -(-len(teks) // chars))
    return lines * pt * 1.25 / 72 + 0.05


def number_pages(prs):
    n = 0
    for slide in prs.slides:
        n += 1
        if slide.slide_id in _PAGED:
            P.add_text(slide, str(n), 12.0, 7.08, 0.78, 0.22, size=9, color=P.GRAY, align=PP_ALIGN.RIGHT)


# ----------------------------------------------------------------- slide
def add_cover(prs, d: dict[str, Any]):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide, MASTER_COVER)
    P.add_text(slide, report_title(d), 0.67, 1.15, 7.75, 1.55, size=28, color=P.DARK, bold=True)
    if project_name(d):
        P.add_text(slide, project_name(d), 0.69, 3.16, 6.8, 0.58, size=15, color=P.INK)
    meta = " | ".join(x for x in [f"No. {d.get('report_no')}" if d.get("report_no") else "",
                                    f"Rev. {d.get('revision')}" if d.get("revision") not in (None, "") else "",
                                    text_of(d.get("date"))] if x)
    if meta:
        P.add_text(slide, meta, 0.69, 4.02, 6.8, 0.35, size=11, color=P.PRIMARY, bold=True)


def add_summary(prs, d: dict[str, Any]):
    slide = content_slide(prs, "Ringkasan Eksekutif", "Hasil Kajian Teknis")
    findings = findings_of(d)
    high = sum(1 for f in findings if priority_level(f.get("priority", "")) == "high")
    n_steps = sum(len(steps_of(f, d)) for f in findings)
    gaps = as_list(d.get("data_gaps"))
    G.add_kpi_row(slide, [
        {"nilai": str(len(findings)), "label": "Item review", "ket": "seluruh disiplin"},
        {"nilai": str(high), "label": "Prioritas tinggi", "ket": "diselesaikan lebih dulu"},
        {"nilai": str(len(gaps)), "label": "Data gap", "ket": "data atau klarifikasi yang diperlukan"},
        {"nilai": str(n_steps), "label": "Langkah penyelesaian", "ket": "tindakan konkret yang diusulkan"},
    ], L, BODY_TOP, CW, h=1.45, size=32)
    msgs = executive_messages(d) or ["Kajian dilakukan terhadap dokumen yang tersedia dengan bukti terpilih dari *knowledge base*."]
    y = P.add_column_head(slide, "Pesan utama", L, 3.15, CW)
    items: list[Any] = [compact(m, 230) for m in msgs[:4]]
    overall = text_of(d.get("overall_status"))
    if overall:
        items.append({"text": f"Rekomendasi kajian: {overall}", "bold": True})
    P.add_bullets(slide, items, L, y, CW, 0.42 * len(items) + 0.2, size=13)
    caption(slide, "Angka kunci dihitung dari daftar item review, *data gap*, dan langkah penyelesaian pada kajian ini.")


def add_agenda(prs):
    slide = content_slide(prs, "Agenda", "Struktur Kajian Teknis")
    items = ["Tujuan, Lingkup, dan Dasar Kajian", "Dokumen yang Dikaji", "Metodologi dan Klasifikasi Komentar",
             "Data Eksisting dan Komentar Review", "Risiko Utama dan Keterkaitan Antardisiplin",
             "*Data Gap*, Klarifikasi, dan Asumsi", "Langkah Penyelesaian", "Kesimpulan dan Langkah Selanjutnya"]
    for i, label in enumerate(items):
        col, row = i % 2, i // 2
        x = L + col * 6.2
        y = BODY_TOP + 0.15 + row * 1.15
        P.add_text(slide, f"{i + 1:02d}", x, y, 0.8, 0.5, size=24, color=P.PRIMARY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        P.add_text(slide, label, x + 0.9, y, 5.0, 0.5, size=15, color=P.INK, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        G._line(slide.shapes, x, y + 0.72, x + 5.8, y + 0.72, P.OUTLINE, 0.75)


def has_control(d: dict[str, Any]) -> bool:
    return any(d.get(k) not in (None, "", []) for k in ("report_no", "revision", "prepared_by", "reviewed_by", "approved_by"))


def add_control(prs, d: dict[str, Any]):
    slide = content_slide(prs, "Pengendalian Dokumen", "Identitas dan Revisi Kajian")
    rows = [[lab, text_of(d.get(k)) or {"text": "", "na": True}] for lab, k in [
        ("No. laporan / kajian", "report_no"), ("Revisi", "revision"), ("Tanggal terbit", "date"),
        ("Disusun oleh", "prepared_by"), ("Diperiksa oleh", "reviewed_by"), ("Disetujui oleh", "approved_by")]]
    T.add_tabel(slide, ["Parameter", "Keterangan"], rows, L, BODY_TOP, 8.0, col_w=[2.6, 5.4], size=13)


def add_scope(prs, d: dict[str, Any]):
    slide = content_slide(prs, "Tujuan dan Lingkup", "Dasar Pelaksanaan Kajian")
    y = P.add_column_head(slide, "Tujuan kajian", L, BODY_TOP, CW)
    P.add_text(slide, compact(d.get("objective") or "Tujuan kajian belum disebutkan.", 300), L, y, CW, 0.8, size=13)
    y2 = BODY_TOP + 1.65
    ya = P.add_column_head(slide, "Lingkup kajian", LEFT_X, y2, LEFT_W - 0.3)
    scope = as_list(d.get("scope") or d.get("coverage")) or ["Lingkup belum disebutkan"]
    P.add_bullets(slide, [compact(x, 160) for x in scope[:7]], LEFT_X, ya, LEFT_W - 0.3, 3.0, size=13)
    yb = P.add_column_head(slide, "Dasar kajian dan keterbatasan", RIGHT_X - 0.4, y2, RIGHT_W + 0.4)
    basis = as_list(d.get("review_basis") or d.get("evidence_basis")) or [
        "Dokumen yang dikaji dan bukti terpilih dari *knowledge base* geothermal"]
    items = [compact(x, 160) for x in basis[:5]] + [{"text": "Keterbatasan: " + compact(x, 150)} for x in as_list(d.get("limitations"))[:2]]
    P.add_bullets(slide, items, RIGHT_X - 0.4, yb, RIGHT_W + 0.4, 3.0, size=13)


def add_documents(prs, d: dict[str, Any]):
    docs = [x for x in as_list(d.get("documents_reviewed")) if isinstance(x, dict)]
    rows = [[text_of(x.get("title")), text_of(x.get("number") or x.get("code")) or {"text": "", "na": True},
             text_of(x.get("revision")) or {"text": "", "na": True}, text_of(x.get("date")) or {"text": "", "na": True},
             text_of(x.get("discipline")), text_of(x.get("scope") or x.get("status"))] for x in docs]
    if not rows:
        rows = [["Dokumen yang dikaji belum dicatat pada data input", "", "", "", "", ""]]
    pages = table_pages(rows, 8)
    for i, chunk in enumerate(pages):
        sub = "Daftar Dokumen" + (f" ({i + 1}/{len(pages)})" if len(pages) > 1 else "")
        slide = content_slide(prs, "Dokumen yang Dikaji", sub)
        T.add_tabel(slide, ["Judul dokumen", "No./kode", "Rev.", "Tanggal", "Disiplin", "Lingkup kajian"],
                    chunk, L, BODY_TOP, CW, col_w=[4.2, 1.8, 0.7, 1.4, 1.5, 2.6], size=TABLE_PT, gaya="data")


def add_methodology(prs, d: dict[str, Any]):
    slide = content_slide(prs, "Metodologi Kajian", "Alur Pemeriksaan dan Klasifikasi")
    custom = [text_of(x) for x in as_list(d.get("methodology")) if text_of(x)]
    if custom:
        langkah = [{"judul": f"Langkah {i + 1}", "sub": compact(t, 110)} for i, t in enumerate(custom[:5])]
    else:
        langkah = [
            {"judul": "Pahami dokumen", "sub": "Identifikasi klaim dan keputusan teknis yang perlu diperiksa"},
            {"judul": "Telusuri KB", "sub": "Ambil bukti yang relevan saja dari *knowledge base*"},
            {"judul": "Uji konsistensi", "sub": "Narasi, tabel, gambar, perhitungan, satuan, dan revisi"},
            {"judul": "Klasifikasi", "sub": "Temuan, *data gap*, dan tindak lanjut yang dapat ditelusuri"},
        ]
    G.add_flow(slide, langkah, L, BODY_TOP + 0.1, CW, h=1.45)
    classes = d.get("comment_classification")
    if isinstance(classes, dict):
        rows = [[k, text_of(v)] for k, v in classes.items()]
        header = ["Klasifikasi", "Arti"]
        col_w = [2.6, 9.6]
    else:
        header = ["Klasifikasi", "Arti", "Perlakuan"]
        col_w = [2.4, 5.4, 4.4]
        rows = [
            ["Prioritas tinggi", "Isu teknis material yang memengaruhi keselamatan, integritas sumur, atau keputusan utama", "Diselesaikan sebelum tahap berikutnya"],
            ["Prioritas sedang", "Perlu perbaikan atau klarifikasi yang memengaruhi desain, biaya, atau jadwal", "Diselesaikan sebelum dinyatakan selesai"],
            ["Prioritas rendah", "Perbaikan minor yang tidak mengubah dasar keputusan", "Dicatat untuk revisi berikutnya"],
            ["*Data gap*", "Bukti belum cukup untuk kesimpulan yang pasti", "Dilengkapi oleh pemilik data"],
        ]
    T.add_tabel(slide, header, rows, L, BODY_TOP + 2.0, CW, col_w=col_w, size=12)
    caption(slide, "Alur di atas menunjukkan urutan pemeriksaan; tabel menjelaskan arti setiap klasifikasi komentar.")


def add_distribution(prs, d: dict[str, Any]):
    findings = findings_of(d)
    if len(findings) < 2:
        return
    disciplines: list[str] = []
    for f in findings:
        key = f.get("discipline") or "Lainnya"
        if key not in disciplines:
            disciplines.append(key)
    levels = [("Tinggi", "high"), ("Sedang", "medium"), ("Rendah", "low")]
    series = [(lab, [sum(1 for f in findings if (f.get("discipline") or "Lainnya") == disc
                         and priority_level(f.get("priority", "")) == lv) for disc in disciplines])
              for lab, lv in levels]
    series = [s for s in series if sum(s[1])] or series[:1]
    slide = content_slide(prs, "Sebaran Temuan", "per Disiplin dan Prioritas")
    y = P.add_column_head(slide, "Jumlah temuan per disiplin", LEFT_X, BODY_TOP, LEFT_W + 0.4)
    h = min(4.9, 0.75 + 0.55 * len(disciplines))
    colors = {"Tinggi": P.DARK, "Sedang": P.PRIMARY, "Rendah": P.TINT}
    G.add_stacked_chart(slide, disciplines, series, LEFT_X, y, LEFT_W + 0.4, h, horizontal=True,
                        colors=[colors[s[0]] for s in series])
    top = sorted(((sum(1 for f in findings if (f.get("discipline") or "Lainnya") == disc), disc) for disc in disciplines), reverse=True)
    ket = [f"Batang menunjukkan jumlah temuan per disiplin, dipisah menurut prioritas.",
           f"Disiplin dengan temuan terbanyak: {top[0][1]} ({top[0][0]} temuan)."]
    P.add_keterangan(slide, ket, RIGHT_X + 0.1, BODY_TOP + 0.6, RIGHT_W - 0.1)


def add_findings_overview(prs, d: dict[str, Any]):
    findings = findings_of(d)
    order = {"high": 0, "medium": 1, "low": 2}
    findings = sorted(findings, key=lambda f: order[priority_level(f.get("priority", ""))])
    rows = []
    for f in findings:
        existing = compact(f.get("existing"), 150) or "Belum dicatat pada data hasil kajian"
        if f.get("location"):
            existing += f" ({compact(f['location'], 45)})"
        rows.append([f.get("id", ""), f.get("discipline", ""), existing,
                     compact(f.get("title") or f.get("finding", ""), 150),
                     f.get("priority", "") or {"text": "", "na": True}])
    if not rows:
        rows = [["—", "—", "Belum ada item review pada data hasil kajian", "—", "—"]]
    pages = table_pages(rows, 5)
    for i, chunk in enumerate(pages):
        sub = "Data Eksisting dan Komentar Review" + (f" ({i + 1}/{len(pages)})" if len(pages) > 1 else "")
        slide = content_slide(prs, "Item Review", sub)
        T.add_tabel(slide, ["ID", "Disiplin", "Data eksisting (lokasi)", "Komentar review", "Prioritas"], chunk,
                    L, BODY_TOP, CW, col_w=[0.8, 1.4, 4.6, 4.4, 1.0], size=TABLE_PT)


def add_finding_detail(prs, f0: dict[str, Any], d: dict[str, Any] | None = None):
    f = normalize_finding(f0)
    title = compact(f.get("title") or f.get("finding") or "Item review", 95)
    slide = content_slide(prs, f.get("id") or "Item", title)
    T.add_tabel(slide, ["Disiplin", "Prioritas", "Tingkat keyakinan", "Lokasi di dokumen"],
                [[f.get("discipline") or {"text": "", "na": True}, f.get("priority") or {"text": "", "na": True},
                  f.get("confidence") or {"text": "", "na": True},
                  compact(f.get("location"), 90) or {"text": "", "na": True}]],
                L, BODY_TOP, CW, col_w=[2.0, 1.6, 2.0, 6.6], size=TABLE_PT)
    steps = steps_of(f, d)
    review = f.get("finding") or ""
    if f.get("comment") and f.get("resolution_steps"):
        review = f"{review} {f['comment']}".strip()
    y = BODY_TOP + 1.05
    for head, body in [("Data eksisting", f.get("existing") or "Belum dicatat pada data hasil kajian."),
                       ("Komentar review", review or "Belum dinyatakan."),
                       ("Dampak bila tidak diselesaikan", f.get("impact") or "Belum dinyatakan.")]:
        y2 = P.add_column_head(slide, head, LEFT_X, y, LEFT_W, size=13)
        teks = compact(body, 330)
        h = text_height(teks, LEFT_W, 12)
        P.add_text(slide, teks, LEFT_X, y2, LEFT_W, h, size=12)
        y = y2 + h + 0.3
    yr = P.add_column_head(slide, "Langkah penyelesaian", RIGHT_X, BODY_TOP + 1.05, RIGHT_W, size=13)
    lines = []
    for i, st in enumerate(steps[:5], 1):
        extra = "; ".join(x for x in (f"oleh {st['pic']}" if st["pic"] else "", f"target {st['due']}" if st["due"] else "",
                                       f"bukti: {st['evidence']}" if st["evidence"] else "") if x)
        lines.append(compact(st['step'], 160) + (f" ({extra})" if extra else ""))
    if lines:
        box = P.add_text(slide, lines, RIGHT_X, yr, RIGHT_W, 2.6, size=12, space_after=6)
        for i, para in enumerate(box.text_frame.paragraphs, 1):
            P._nomor(para, i)  # penomoran otomatis PowerPoint
    else:
        lines = ["Tidak ada langkah yang diperlukan; data eksisting sudah memadai." if tanpa_tindakan(f)
                 else "Langkah penyelesaian belum dirumuskan."]
        P.add_text(slide, lines, RIGHT_X, yr, RIGHT_W, 0.4, size=12)
    yb = yr + sum(text_height(t, RIGHT_W - 0.3, 12) + 0.08 for t in lines) + 0.25
    yb = P.add_column_head(slide, "Dasar teknis", RIGHT_X, yb, RIGHT_W, size=13)
    img = f.get("image_path")
    if img and os.path.exists(str(img)):
        T.fit_picture(slide, str(img), RIGHT_X, yb, RIGHT_W, 1.6)
        yb += 1.7
    P.add_text(slide, compact(f.get("basis") or "Belum dinyatakan.", 220), RIGHT_X, yb, RIGHT_W, 0.9, size=11)
    add_source(slide, ("Sumber: " + f["source"]) if f.get("source") else "")


def add_risks(prs, d: dict[str, Any]):
    risks = as_list(d.get("key_risks"))
    slide = content_slide(prs, "Risiko Utama", "Keterkaitan Antardisiplin")
    if not risks:
        P.add_text(slide, "Belum ada risiko keterkaitan antardisiplin yang dicatat pada data hasil kajian.",
                   L, BODY_TOP, CW, 0.5, size=13)
        return
    dict_risks = [r for r in risks if isinstance(r, dict)]
    scored = [r for r in dict_risks if str(r.get("likelihood") or r.get("kemungkinan") or "").isdigit()
              and str(r.get("severity") or r.get("dampak_level") or "").isdigit()]
    rows = []
    for i, r in enumerate(risks[:6], 1):
        if isinstance(r, dict):
            rows.append([str(i), compact(clean(r.get("interface") or r.get("title") or "—"), 60),
                         compact(clean(r.get("risk") or r.get("issue") or ""), 140),
                         compact(r.get("impact") or r.get("consequence") or "", 110) or {"text": "", "na": True},
                         compact(r.get("mitigation") or r.get("mitigasi") or "", 120) or {"text": "", "na": True},
                         text_of(r.get("related_findings") or r.get("finding_id")) or {"text": "", "na": True}])
        else:
            rows.append([str(i), "—", compact(r, 160), {"text": "", "na": True}, {"text": "", "na": True}, {"text": "", "na": True}])
    if scored:
        sebaran: dict[tuple[int, int], list[int]] = {}
        for i, r in enumerate(risks[:6], 1):
            if r in scored:
                k = int(r.get("likelihood") or r.get("kemungkinan"))
                s = int(r.get("severity") or r.get("dampak_level"))
                sebaran.setdefault((max(1, min(5, k)), max(1, min(5, s))), []).append(i)
        T.add_matriks_risiko(slide, sebaran, L, BODY_TOP, 5.2, 5.2, judul="Peta risiko (nomor = baris tabel)")
        T.add_tabel(slide, ["No.", "Keterkaitan", "Risiko", "Mitigasi"], [[r[0], r[1], r[2], r[4]] for r in rows],
                    6.0, BODY_TOP, R - 6.0, col_w=[0.5, 1.5, 2.6, 2.2], size=10)
        caption(slide, "Nomor pada peta merujuk baris tabel; warna sel menunjukkan tingkat risiko dari rendah (hijau) "
                       "hingga tinggi (merah).", x=6.0, w=R - 6.0)
    else:
        T.add_tabel(slide, ["No.", "Keterkaitan", "Risiko", "Dampak", "Mitigasi", "Temuan terkait"], rows,
                    L, BODY_TOP, CW, col_w=[0.5, 1.9, 3.4, 2.5, 2.8, 1.1], size=TABLE_PT)


def gap_steps(g: dict[str, Any]) -> list[str] | str:
    """Langkah melengkapi data gap: resolution_steps/action bila ada, selain itu data yang diminta."""
    steps = [text_of(x.get("step") if isinstance(x, dict) else x)
             for x in as_list(g.get("resolution_steps") or g.get("steps") or g.get("action"))]
    steps = [x for x in steps if x]
    if steps:
        return numbered(steps[:4], 110) if len(steps) > 1 else compact(steps[0], 130)
    req = text_of(g.get("required") or g.get("evidence"))
    return compact(f"Sediakan {req}", 130) if req else ""


def add_gaps(prs, d: dict[str, Any]):
    rows = []
    for i, g in enumerate(as_list(d.get("data_gaps"))):
        if isinstance(g, dict):
            rows.append([text_of(g.get("id") or f"DG-{i + 1:02d}"), compact(g.get("gap") or g.get("item"), 140),
                         compact(g.get("impact"), 110) or {"text": "", "na": True},
                         gap_steps(g) or {"text": "", "na": True},
                         text_of(g.get("owner") or g.get("pic")) or {"text": "", "na": True}])
        else:
            rows.append([f"DG-{i + 1:02d}", compact(g, 140), {"text": "", "na": True}, {"text": "", "na": True},
                         {"text": "", "na": True}])
    if not rows:
        rows = [["—", "Tidak ada *data gap* material yang tercatat", "—", "—", "—"]]
    pages = table_pages(rows, 6)
    for i, chunk in enumerate(pages):
        sub = "Data yang Perlu Dilengkapi" + (f" ({i + 1}/{len(pages)})" if len(pages) > 1 else "")
        slide = content_slide(prs, "*Data Gap*", sub)
        T.add_tabel(slide, ["ID", "Data yang belum tersedia / belum jelas", "Dampak bila tidak dipenuhi",
                            "Langkah untuk melengkapi", "PIC"], chunk, L, BODY_TOP, CW,
                    col_w=[0.8, 3.7, 2.8, 3.6, 1.3], size=TABLE_PT)


def add_actions(prs, d: dict[str, Any]):
    """Langkah penyelesaian per item review (menggantikan register status tindak lanjut)."""
    findings = findings_of(d)
    order = {"high": 0, "medium": 1, "low": 2}
    findings = sorted(findings, key=lambda f: order[priority_level(f.get("priority", ""))])
    rows = []
    for f in findings:
        steps = steps_of(f, d)
        if not steps:
            continue
        pics = sorted({s["pic"] for s in steps if s["pic"]})
        dues = [s["due"] for s in steps if s["due"]]
        evid = [s["evidence"] for s in steps if s["evidence"]]
        rows.append([f.get("id", ""), compact(f.get("title") or f.get("finding"), 90),
                     numbered([s["step"] for s in steps[:5]], 130),
                     ", ".join(pics) or {"text": "", "na": True}, "; ".join(dues) or {"text": "", "na": True},
                     compact("; ".join(evid), 120) or {"text": "", "na": True}])
    linked = {f.get("id") for f in findings}
    for i, a in enumerate(as_list(d.get("actions"))):
        ref = text_of(a.get("finding_id") or a.get("related_findings")) if isinstance(a, dict) else ""
        if isinstance(a, dict) and any(fid and fid in ref for fid in linked):
            continue
        a = a if isinstance(a, dict) else {"action": text_of(a)}
        rows.append([text_of(a.get("id") or f"A-{i + 1:02d}"), compact(ref, 60) or {"text": "", "na": True},
                     numbered([text_of(a.get("action"))], 130), text_of(a.get("pic") or a.get("owner")) or {"text": "", "na": True},
                     text_of(a.get("due") or a.get("timing")) or {"text": "", "na": True},
                     compact(a.get("evidence") or a.get("closure_evidence"), 120) or {"text": "", "na": True}])
    if not rows:
        rows = [["—", "—", ["1. Rumuskan langkah penyelesaian untuk setiap item review dan *data gap*."], "—", "—", "—"]]
    pages = table_pages(rows, 4)
    for i, chunk in enumerate(pages):
        sub = "Langkah Konkret per Item Review" + (f" ({i + 1}/{len(pages)})" if len(pages) > 1 else "")
        slide = content_slide(prs, "Langkah Penyelesaian", sub)
        T.add_tabel(slide, ["ID", "Item review", "Langkah penyelesaian", "PIC", "Target", "Bukti penyelesaian"],
                    chunk, L, BODY_TOP, CW, col_w=[0.7, 2.4, 5.0, 1.5, 1.1, 1.8], size=10.5)


def add_conclusion(prs, d: dict[str, Any]):
    slide = content_slide(prs, "Kesimpulan", "Langkah Selanjutnya")
    y = P.add_column_head(slide, "Kesimpulan kajian", LEFT_X, BODY_TOP, LEFT_W)
    teks = compact(d.get("conclusion") or "Kesimpulan belum disertakan.", 520)
    h = text_height(teks, LEFT_W, 13)
    P.add_text(slide, teks, LEFT_X, y, LEFT_W, h, size=13)
    overall = text_of(d.get("overall_status"))
    if overall:
        P.add_text(slide, f"Rekomendasi kajian: **{overall}**", LEFT_X, y + h + 0.25, LEFT_W, 0.5, size=13)
    nxt = [text_of(x) for x in as_list(d.get("next_steps") or d.get("conditions_to_proceed") or d.get("closeout_conditions")) if text_of(x)]
    if not nxt:  # turunkan dari langkah pertama tiap item prioritas tertinggi
        order = {"high": 0, "medium": 1, "low": 2}
        for f in sorted(findings_of(d), key=lambda f: order[priority_level(f.get("priority", ""))]):
            st = steps_of(f, d)
            if st:
                nxt.append(f"{st[0]['step']} ({f.get('id', '')})")
            if len(nxt) == 5:
                break
    yr = P.add_column_head(slide, "Langkah selanjutnya (berurutan)", RIGHT_X - 0.2, BODY_TOP, RIGHT_W + 0.2)
    for i, t in enumerate(nxt[:6], 1):
        body = compact(t, 170)
        hh = max(0.45, text_height(body, RIGHT_W - 0.45, 12))
        P.add_text(slide, f"{i}", RIGHT_X - 0.2, yr, 0.4, 0.35, size=16, color=P.PRIMARY, bold=True)
        P.add_text(slide, body, RIGHT_X + 0.25, yr + 0.03, RIGHT_W - 0.45, hh, size=12)
        yr += hh + 0.15
    if not nxt:
        P.add_text(slide, "Langkah selanjutnya belum dirumuskan.", RIGHT_X - 0.2, yr, RIGHT_W, 0.4, size=12)


def add_references(prs, d: dict[str, Any]):
    refs = [text_of(r) for r in as_list(d.get("references_used") or d.get("references")) if text_of(r)]
    if not refs:
        refs = sorted({f.get("source", "") for f in findings_of(d) if f.get("source")})
    slide = content_slide(prs, "Lampiran", "Sumber yang Digunakan")
    if refs:
        rows = [[f"{i:02d}", compact(r, 200)] for i, r in enumerate(refs[:12], 1)]
        T.add_tabel(slide, ["No.", "Sumber (file dan halaman/slide)"], rows, L, BODY_TOP, CW, col_w=[0.7, 11.5],
                    size=TABLE_PT, sumbu=False)
    else:
        P.add_text(slide, "Daftar sumber belum disertakan pada data hasil kajian.", L, BODY_TOP, CW, 0.4, size=13)
    add_source(slide, "Hanya sumber yang benar-benar dibuka dan dipakai sebagai bukti teknis yang dicantumkan.")


def add_closing(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide, MASTER_CLOSING)


def build(data: dict[str, Any], out_path: str, extra=None):
    """extra: fungsi opsional extra(prs, data) untuk menyisipkan slide tambahan
    (gantt, stage-gate, peta, dll.) sebelum slide kesimpulan."""
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    add_cover(prs, data)
    add_summary(prs, data)
    if data.get("include_agenda", True):
        add_agenda(prs)
    if has_control(data) and data.get("include_review_control", True):
        add_control(prs, data)
    add_scope(prs, data)
    add_documents(prs, data)
    add_methodology(prs, data)
    add_distribution(prs, data)
    add_findings_overview(prs, data)
    detail_limit = int(data.get("ppt_detail_limit", 6) or 6)
    for f in [x for x in as_list(data.get("findings")) if isinstance(x, dict)][:detail_limit]:
        add_finding_detail(prs, f, data)
    add_risks(prs, data)
    add_gaps(prs, data)
    add_actions(prs, data)
    if extra:
        extra(prs, data)
    add_conclusion(prs, data)
    add_references(prs, data)
    add_closing(prs)
    number_pages(prs)
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)


def main() -> int:
    if len(sys.argv) != 3:
        print(__doc__)
        return 2
    with open(sys.argv[1], "r", encoding="utf-8") as f:
        data = json.load(f)
    build(data, sys.argv[2])
    print(sys.argv[2])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
