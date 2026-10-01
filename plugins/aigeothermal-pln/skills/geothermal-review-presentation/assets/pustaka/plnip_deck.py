"""plnip_deck.py — pembuat deck PLN IP dari MASTER NATIVE korporat.

Dasar deck (default): assets/masters/PLN_IP_Master_Library.pptx — 38 desain master asli
(cover berilustrasi, cover foto, pembatas bab, isi, penutup). Logo, latar, dan artwork
tetap berada di slide master; slide hanya berisi konten. Master yang tidak dipakai dibuang
otomatis saat save(), jadi deck akhir hanya membawa desain yang benar-benar dipakai.
Jalur lama (Template_PLNIP.pptx) tetap ada: new_deck(dasar='template').

Pakai:
    import sys; sys.path.insert(0, '<folder assets>')
    from plnip_deck import *

    prs = new_deck()                                   # desain isi default = 1
    set_cover(prs, 'Kajian Pemanfaatan Lahan Jambi', 'Divisi GRB', 'September 2026', desain=12)
    add_divider(prs, 'Latar Belakang', nomor=1)        # pembatas bab (desain 20)
    s = add_slide(prs, 'Latar Belakang', 'Proyek PLTMG Minahasa',
)
    add_bullets(s, ['...'], LEFT_X, BODY_TOP, LEFT_W, 3.0)          # fakta/data di kiri
    add_keterangan(s, 'Status pembebasan lahan per bidang, Agustus 2026.',
                   RIGHT_X, BODY_TOP, RIGHT_W)           # keterangan singkat di badan slide
    save(prs, 'keluaran.pptx')

Model grafis (waterfall, gantt, peta Indonesia, swimlane, matriks 2x2, dll.) ada di
plnip_grafis.py dan ikut ter-impor lewat `from plnip_deck import *`.
Semua ukuran dalam inci. Kanvas 13,333 x 7,5 (16:9).
"""

import copy
import os
import re

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, 'example', 'Template_PLNIP.pptx')
MASTER_LIB = os.path.join(HERE, 'masters', 'PLN_IP_Master_Library.pptx')

# ---------- Palet (lihat SKILL.md) ----------
A1 = RGBColor(0x1B, 0x4F, 0x60)
A2 = PRIMARY = RGBColor(0x00, 0x8A, 0xAC)
A3 = DARK = RGBColor(0x05, 0x36, 0x5B)
A4 = RGBColor(0x29, 0x53, 0x7B)
A5 = TINT = RGBColor(0xD1, 0xED, 0xF3)
YELLOW = RGBColor(0xFF, 0xD9, 0x66)
TEAL = RGBColor(0x0D, 0xAD, 0x8E)
BLUE_LAMA = RGBColor(0x00, 0x66, 0x99)
MIST = RGBColor(0xF6, 0xF9, 0xFB)
INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x4D, 0x4D, 0x4D)
OUTLINE = RGBColor(0xC0, 0xC0, 0xC0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TL_GREEN = RGBColor(0x22, 0xB4, 0x4A)
TL_YELLOW = RGBColor(0xEF, 0xCF, 0x06)
TL_RED = RGBColor(0xE0, 0x01, 0x02)
NEG = RGBColor(0xB0, 0x3A, 0x2E)

TOPIC_COLOR = PRIMARY   # warna bagian topik pada judul
SUBTOPIC_COLOR = INK    # warna bagian sub-topik pada judul

FONT = 'Helvetica'      # font korporat wajib (aturan identitas B). LibreOffice merender dengan Liberation Sans

# ---------- Grid (diukur dari template resmi) ----------
W, H = 13.333, 7.5
L, R = 0.5, 12.833
CW = R - L
TITLE_Y, TITLE_H, TITLE_PT = 0.28, 0.55, 28   # Helvetica lebih lebar dari Calibri: 28pt agar 2 baris cukup
TITLE_W = 8.61                                 # berhenti sebelum logo di kanan atas
BODY_TOP, BODY_BOTTOM = 1.3, 6.85             # tanpa narasi, baris sumber, dan footer teks
KET_PT = 13                                    # ukuran teks keterangan di badan slide
PESAN_PT = KET_PT
HEAD_H = 0.42                                  # tinggi judul kolom / header tabel
LEFT_X, LEFT_W, RIGHT_X, RIGHT_W = 0.5, 6.6, 7.5, 5.333
HALF_W, HALF2_X = 5.966, 6.866


# ---------- Dasar ----------
# ---------- Katalog desain master (ID = assets/masters/catalog.json) ----------
# peran : 'isi' (slide isi standar: judul kiri atas, logo kanan atas), 'isi-foto' (foto di kiri,
#         isi mulai x0), 'cover', 'pembatas', 'penutup', 'gambar' (foto penuh tanpa teks),
#         'khusus' (tata letak header berbeda; hanya bila diminta).
# zona  : kotak teks (x, y, w, h) untuk cover/pembatas/penutup. terang=True -> teks putih.
# nomor : False bila master sudah membawa nomor halaman sendiri.
DESAIN = {
    1:  dict(nama='Isi putih, siluet pembangkit', peran='isi'),
    2:  dict(nama='Ilustrasi PLTS', peran='cover', zona=(7.7, 2.2, 5.1, 3.2)),
    3:  dict(nama='Ilustrasi PLTU', peran='cover', zona=(7.7, 2.2, 5.1, 3.2)),
    4:  dict(nama='Ilustrasi panas bumi', peran='cover', zona=(7.7, 2.2, 5.1, 3.2)),
    5:  dict(nama='Ilustrasi pipa hidro', peran='cover', zona=(7.9, 2.2, 4.9, 3.2)),
    6:  dict(nama='Ilustrasi pembangkit gas', peran='cover', zona=(7.9, 2.2, 4.9, 3.2)),
    7:  dict(nama='Ilustrasi bendungan', peran='cover', zona=(7.9, 2.2, 4.9, 3.2)),
    8:  dict(nama='Ilustrasi panas bumi hijau', peran='cover', zona=(7.9, 2.2, 4.9, 3.2)),
    9:  dict(nama='Isi latar PLTS lembut', peran='isi', nomor=False),
    10: dict(nama='Foto PLTS dan gunung', peran='isi-foto', x0=4.95),
    11: dict(nama='Foto panel surya', peran='isi-foto', x0=4.95),
    12: dict(nama='Cover energi hijau', peran='cover', zona=(8.0, 1.6, 4.9, 3.6)),
    13: dict(nama='Isi langit biru', peran='isi'),
    14: dict(nama='Foto barisan petugas PLN', peran='gambar'),
    15: dict(nama='Cover alam dan komunitas', peran='cover', zona=(7.2, 2.0, 5.6, 3.0)),
    16: dict(nama='Cover pekerja dan energi', peran='cover', zona=(0.67, 1.3, 7.8, 2.8)),
    17: dict(nama='Biru jaringan transmisi', peran='pembatas', zona=(0.6, 2.2, 8.0, 3.0), terang=True),
    18: dict(nama='Biru pembangkit malam', peran='pembatas', zona=(0.6, 2.2, 7.5, 3.0), terang=True),
    19: dict(nama='Cover biru chevron', peran='cover', zona=(0.55, 1.4, 6.2, 3.4), terang=True),
    20: dict(nama='Divider biru chevron', peran='pembatas', zona=(0.8, 2.0, 10.0, 3.0), terang=True),
    21: dict(nama='Isi gradasi biru muda', peran='khusus'),
    22: dict(nama='Penutup biru chevron', peran='penutup', zona=(0.5, 2.3, 10.0, 2.6), terang=True),
    23: dict(nama='Isi putih logo PLN', peran='isi'),
    24: dict(nama='Isi gelombang biru', peran='isi'),
    25: dict(nama='Cover PLTS terapung', peran='cover', zona=(6.5, 2.0, 6.4, 2.9), terang=True),
    26: dict(nama='Cover PLTP Kamojang', peran='cover', zona=(6.5, 2.0, 6.4, 2.9), terang=True),
    27: dict(nama='Cover PLTA', peran='cover', zona=(6.5, 2.0, 6.4, 2.9), terang=True),
    28: dict(nama='Isi putih PLN IP confidential', peran='isi', nomor=False),
    29: dict(nama='Divider pita teal', peran='pembatas', zona=(0.6, 2.2, 9.0, 3.0), terang=True),
    30: dict(nama='Cover panel PLTS', peran='cover', zona=(5.7, 2.2, 6.0, 2.7), terang=True),
    31: dict(nama='Cover fasilitas pembangkit', peran='cover', zona=(5.7, 2.2, 6.0, 2.7), terang=True),
    32: dict(nama='Cover PLTA Bengkok', peran='cover', zona=(5.7, 2.2, 6.0, 2.7), terang=True),
    33: dict(nama='Isi footer teal', peran='khusus'),
    34: dict(nama='Divider blok teal kiri', peran='pembatas', zona=(6.1, 2.3, 6.6, 2.8), blok=(0.5, 2.3, 4.6, 2.8)),
    35: dict(nama='Cover teknologi digital', peran='cover', zona=(0.6, 2.3, 6.5, 2.8), terang=True),
    36: dict(nama='Isi legacy BUMN', peran='khusus'),
    37: dict(nama='Divider PLN IP Renewables', peran='pembatas', zona=(0.8, 2.2, 6.0, 3.0), terang=True),
    38: dict(nama='Isi PLN IP Renewables', peran='isi'),
}
# Tepi kiri blok logo kanan atas per desain isi (diukur dari render master). Judul berhenti 0,25 in sebelumnya.
LOGO_KIRI = {1: 9.36, 9: 10.76, 10: 9.36, 11: 9.36, 13: 9.35, 23: 11.16, 24: 9.35, 28: 10.84, 38: 11.05}

_STATE = {'dasar': 'master', 'isi': 1}


def new_deck(dasar='master', desain_isi=1, template=None, keep_cover=True):
    """Deck baru.

    dasar='master'   (default) pustaka 38 master native; slide ditambahkan dengan add_slide /
                     set_cover / add_divider / add_closing, masing-masing memilih desainnya.
    dasar='template' jalur lama: membuka Template_PLNIP.pptx (cover berilustrasi + satu master).
    desain_isi       desain default untuk add_slide (1 = isi putih siluet pembangkit).
    """
    if dasar == 'template' or template:
        path = template or TEMPLATE
        if not os.path.exists(path):
            raise FileNotFoundError('Template tidak ketemu: ' + path)
        _STATE.update(dasar='template')
        prs = Presentation(path)
        _drop_slide(prs, 1)                      # slide "Judul" kosong bawaan template
        if not keep_cover:
            _drop_slide(prs, 0)
        return prs
    if not os.path.exists(MASTER_LIB):
        raise FileNotFoundError(
            'Pustaka master tidak ketemu: ' + MASTER_LIB +
            '. Salin seluruh folder assets/ ke direktori kerja (termasuk assets/masters/).')
    if DESAIN.get(desain_isi, {}).get('peran') not in ('isi', 'isi-foto', 'khusus'):
        raise ValueError('desain_isi %s bukan desain isi' % desain_isi)
    _STATE.update(dasar='master', isi=desain_isi)
    prs = Presentation(MASTER_LIB)
    for _ in range(len(prs.slides)):
        _drop_slide(prs, 0)
    return prs


def _layout(prs, desain):
    """Layout tunggal milik master desain n (dicocokkan lewat nama 'PLN IP nn - ...')."""
    kode = 'PLN IP %02d' % desain
    for m in prs.slide_masters:
        for lay in m.slide_layouts:
            if lay.name.startswith(kode):
                return lay
    raise ValueError('Desain %s tidak ada di pustaka (1-38) atau sudah dibuang' % desain)


def _new_slide(prs, desain):
    """Slide kosong dari master desain n. Placeholder bawaan layout dihapus agar tidak ada
    kotak 'Klik untuk menambah teks' yang tertinggal; artwork master tetap tampil."""
    s = prs.slides.add_slide(_layout(prs, desain))
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    s._plnip_desain = desain
    return s


def _move_slide(prs, slide, index):
    lst = prs.slides._sldIdLst
    for el in list(lst):
        if prs.part.related_part(el.rId) is slide.part:
            lst.remove(el)
            lst.insert(index, el)
            return


def _prune_masters(prs):
    """Buang master (beserta layout, tema, dan gambarnya) yang tidak dipakai slide mana pun."""
    used = {s.slide_layout.slide_master.part for s in prs.slides}
    lst = prs.slide_masters._sldMasterIdLst
    for el in list(lst):
        part = prs.part.related_part(el.rId)
        if part not in used:
            rId = el.rId
            lst.remove(el)
            prs.part.drop_rel(rId)


def _drop_slide(prs, index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    if index < len(slides):
        rId = slides[index].get(qn('r:id'))
        prs.part.drop_rel(rId)
        xml_slides.remove(slides[index])


def save(prs, path):
    """Menyimpan deck. Bagian revisionInfo bawaan template (khas think-cell/OneDrive) dibuang
    supaya berkasnya lolos validator OOXML; isinya tidak memengaruhi tampilan."""
    if _STATE['dasar'] == 'master':
        _prune_masters(prs)
    prs.save(path)
    _strip_part(path, 'ppt/revisionInfo.xml')
    return path


def _strip_part(path, part_name):
    import shutil
    import zipfile
    src = zipfile.ZipFile(path)
    names = src.namelist()
    if part_name not in names:
        src.close()
        return
    tmp = path + '.tmp'
    with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            if item.filename == part_name:
                continue
            data = src.read(item.filename)
            if item.filename == '[Content_Types].xml':
                data = re.sub(rb'<Override PartName="/' + re.escape(part_name.encode()) +
                              rb'"[^>]*/>', b'', data)
            if item.filename == 'ppt/_rels/presentation.xml.rels':
                data = re.sub(rb'<Relationship[^>]*revisionInfo\.xml"[^>]*/>', b'', data)
            dst.writestr(item, data)
    src.close()
    shutil.move(tmp, path)


def _blank_layout(prs):
    for lay in prs.slide_layouts:
        if lay.name.lower().startswith('blank'):
            return lay
    return prs.slide_layouts[0]


def set_cover(prs, judul, unit='', tanggal='', desain=12, subjudul=''):
    """Cover deck.

    Jalur master: membuat slide cover dari desain `desain` (lihat DESAIN, peran 'cover') dan
    menaruhnya sebagai slide pertama. Pilih desain yang bertema dengan isi deck:
    2 PLTS, 3 PLTU, 4/8 panas bumi, 5/7 hidro, 6 gas, 12 energi hijau (default), 25 PLTS terapung,
    26 PLTP, 27/32 PLTA, 30 panel PLTS, 31 fasilitas pembangkit, 19/35 biru/digital, 15/16 komunitas.
    Jalur template: mengisi cover bawaan Template_PLNIP.pptx.
    """
    if _STATE['dasar'] == 'template':
        if not prs.slides:
            return None
        s = prs.slides[0]
        teks = [sh for sh in s.shapes if sh.has_text_frame]
        if teks:
            _set_text(teks[0], judul)
        if len(teks) > 1:
            baris = [x for x in (unit, tanggal) if x]
            _set_text(teks[1], '\n'.join(baris))
        return s
    prof = DESAIN.get(desain, {})
    if prof.get('peran') != 'cover':
        raise ValueError('Desain %s bukan cover. Cover: %s' % (
            desain, [k for k, v in DESAIN.items() if v['peran'] == 'cover']))
    s = _new_slide(prs, desain)
    x, y, w, h = prof['zona']
    warna = WHITE if prof.get('terang') else DARK
    warna2 = WHITE if prof.get('terang') else GRAY
    size = 32 if len(judul) <= 45 else 28
    box = add_text(s, judul, x, y, w, h * 0.62, size=size, color=warna, bold=True,
                   anchor=MSO_ANCHOR.BOTTOM, line=0.95)
    box.name = 'Title'
    yy = y + h * 0.62 + 0.15
    if subjudul:
        add_text(s, subjudul, x, yy, w, 0.4, size=16, color=warna)
        yy += 0.5
    meta = '  |  '.join(t for t in (unit, tanggal) if t)
    if meta:
        add_text(s, meta, x, yy, w, 0.35, size=13, color=warna2)
    _move_slide(prs, s, 0)
    return s


def add_divider(prs, judul, nomor=None, keterangan='', desain=20):
    """Pembatas bab bergaya master (desain peran 'pembatas': 17, 18, 20, 29, 34, 37).
    Judul slide otomatis 'Bab n' + judul bab agar terbaca checker sebagai pembatas.
    Alternatif yang lebih ringan: add_section_slide (content tracker di slide isi)."""
    if _STATE['dasar'] == 'template':
        raise RuntimeError('add_divider hanya untuk jalur master; pakai add_section_slide')
    prof = DESAIN.get(desain, {})
    if prof.get('peran') != 'pembatas':
        raise ValueError('Desain %s bukan pembatas. Pembatas: %s' % (
            desain, [k for k, v in DESAIN.items() if v['peran'] == 'pembatas']))
    s = _new_slide(prs, desain)
    x, y, w, h = prof['zona']
    warna = WHITE if prof.get('terang') else DARK
    if prof.get('blok') and nomor is not None:        # desain 34: nomor besar di blok teal
        bx, by, bw, bh = prof['blok']
        add_text(s, '%02d' % nomor, bx, by, bw, bh, size=96, color=WHITE, bold=True,
                 anchor=MSO_ANCHOR.MIDDLE)
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h * 0.6))
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.BOTTOM
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    if nomor is not None:                  # 'Bab n' kecil di atas judul; dibaca checker sbg pembatas
        _run(p, 'Bab %d' % nomor, 18, warna)
        p.add_line_break()
    _fill_runs(p, judul, 40, warna, bold=True)
    box.name = 'Judul pembatas'
    if keterangan:
        add_text(s, keterangan, x, y + h * 0.6 + 0.2, w, h * 0.4 - 0.2, size=14,
                 color=warna if prof.get('terang') else GRAY)
    return s


def add_closing(prs, judul='PT PLN Indonesia Power', kontak=(), desain=22):
    """Slide penutup (desain 22 biru chevron, atau desain cover lain). Tidak membawa pesan;
    ajakan keputusan ditaruh di slide isi sebelumnya. kontak: baris nama/unit/email."""
    prof = DESAIN.get(desain, {})
    if prof.get('peran') not in ('penutup', 'cover', 'pembatas'):
        raise ValueError('Desain %s tidak cocok untuk penutup' % desain)
    s = _new_slide(prs, desain)
    x, y, w, h = prof['zona']
    warna = WHITE if prof.get('terang') else DARK
    box = add_text(s, judul, x, y, w, 0.9, size=32, color=warna, bold=True, anchor=MSO_ANCHOR.BOTTOM)
    box.name = 'Judul penutup'
    if kontak:
        add_text(s, list(kontak), x, y + 1.05, w, h - 1.05, size=14, color=warna)
    return s


def add_image_slide(prs, desain=14):
    """Slide gambar penuh tanpa teks (desain 14). Pakai hemat: pembuka/penutup acara."""
    return _new_slide(prs, desain)


def _set_text(shape, teks):
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    font_ref = p0.runs[0].font if p0.runs else None
    size = font_ref.size if font_ref is not None else None
    color = None
    if font_ref is not None and font_ref.color and font_ref.color.type is not None:
        try:
            color = font_ref.color.rgb
        except AttributeError:
            color = None
    bold = font_ref.bold if font_ref is not None else None
    align = p0.alignment
    tf.clear()
    for i, baris in enumerate(str(teks).split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        _no_bullet(p)
        r = p.add_run(); r.text = baris
        if size: r.font.size = size
        if bold is not None: r.font.bold = bold
        if color is not None: r.font.color.rgb = color


def _no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    for tag in ('a:buChar', 'a:buAutoNum', 'a:buNone'):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    pPr.append(pPr.makeelement(qn('a:buNone'), {}))


# ---------- Teks ----------
_ITALIC = re.compile(r'\*\*(.+?)\*\*|\*(.+?)\*')


def _fill_runs(p, teks, size, color, bold=False, font=FONT):
    """Menulis teks ke paragraf. Bagian di antara *bintang* menjadi miring
    (dipakai untuk istilah Inggris: 'skema *take or pay* berlaku ...').
    Tambahan AIGeothermal-PLN: **dua bintang** menjadi tebal."""
    pos = 0
    for m in _ITALIC.finditer(teks):
        if m.start() > pos:
            _run(p, teks[pos:m.start()], size, color, bold, font)
        if m.group(1) is not None:
            _run(p, m.group(1), size, color, True, font)
        else:
            _run(p, m.group(2), size, color, bold, font, italic=True)
        pos = m.end()
    if pos < len(teks):
        _run(p, teks[pos:], size, color, bold, font)


def _run(p, teks, size, color, bold=False, font=FONT, italic=False):
    r = p.add_run(); r.text = teks
    r.font.size = Pt(size); r.font.name = font; r.font.bold = bold
    r.font.italic = italic
    if color is not None:
        r.font.color.rgb = color
    return r


def add_text(slide, teks, x, y, w, h, size=13, color=INK, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=4, line=1.0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    baris = teks if isinstance(teks, (list, tuple)) else [teks]
    for i, b in enumerate(baris):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if line != 1.0:
            p.line_spacing = line
        _fill_runs(p, str(b), size, color, bold)
    return box


# ---------- Slide isi ----------
def set_footer(label=None, deck=None):
    """Tidak dipakai lagi: slide isi tidak memakai footer teks ('PT PLN Indonesia Power | ...').
    Dibiarkan ada agar kode lama tidak error; tidak menulis apa pun."""
    return None


def add_slide(prs, topik, subtopik='', narasi=None, source=None, desain=None):
    """Slide isi baru: judul dua warna + nomor halaman. Tidak ada teks lain di luar badan slide.

    topik    : label topik, dicetak BIRU  ('Latar Belakang', 'Analisis Biaya', ...)
    subtopik : lanjutan judul, dicetak HITAM ('Proyek PLTMG Minahasa')
    desain   : desain master isi (default dari new_deck). 1, 9, 13, 23, 24, 28, 38 = isi standar;
               10/11 = foto di kiri (isi mulai FOTO_L = 4,95 in).

    Tidak ada teks kecil di bawah judul (khas deck buatan AI). Keterangan singkat tentang grafis
    ditulis di badan slide: kolom kanan/kiri dengan add_keterangan(), atau di bawah konten dengan
    add_keterangan_bawah(). Baris 'Sumber: ...' dan footer teks juga tidak dipakai (materi internal).
    """
    if narasi is not None:
        raise ValueError('narasi di bawah judul tidak dipakai lagi. Tulis keterangan singkat grafis di '
                         'badan slide: add_keterangan(s, teks, RIGHT_X, BODY_TOP, RIGHT_W) atau '
                         'add_keterangan_bawah(s, teks).')
    if source is not None:
        raise ValueError('Baris sumber tidak dipakai (materi internal). Hapus argumen source=; '
                         'bila asumsi perlu disebut, tulis di badan slide sebagai bagian isi.')
    x0, tw, nomor = L, TITLE_W, True
    if _STATE['dasar'] == 'master':
        d = desain or _STATE['isi']
        prof = DESAIN.get(d, {})
        if prof.get('peran') not in ('isi', 'isi-foto', 'khusus'):
            raise ValueError('Desain %s bukan desain isi (peran %s)' % (d, prof.get('peran')))
        s = _new_slide(prs, d)
        x0 = prof.get('x0', L)
        tw = min(LOGO_KIRI.get(d, 9.36) - 0.25 - x0, 9.3)
        nomor = prof.get('nomor', True)
    else:
        s = prs.slides.add_slide(_blank_layout(prs))
    # Judul dimuatkan: 28pt -> 26 -> 24 bila perlu agar tetap satu baris (persingkat sub-topik bila
    # muncul PERINGATAN).
    full = (topik + ' ' + subtopik).replace('*', '').strip()
    pt, n_lines = TITLE_PT, _title_lines(full, tw, TITLE_PT)
    for cand in (TITLE_PT - 2, TITLE_PT - 4):
        if n_lines == 1:
            break
        pt, n_lines = cand, _title_lines(full, tw, cand)
    if n_lines > 1:
        print('PERINGATAN S%d: judul "%s" butuh %d baris di lebar %.2f in; persingkat sub-topik.'
              % (len(prs.slides._sldIdLst), full, n_lines, tw))
    title_h = max(TITLE_H, n_lines * pt * 1.18 / 72)
    box = s.shapes.add_textbox(Inches(x0), Inches(TITLE_Y), Inches(tw), Inches(title_h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    _fill_runs(p, topik, pt, TOPIC_COLOR, bold=True)
    if subtopik:
        _run(p, ' ', pt, SUBTOPIC_COLOR, bold=True)
        _fill_runs(p, subtopik, pt, SUBTOPIC_COLOR, bold=True)
    box.name = 'Title'
    if nomor:
        add_text(s, str(len(prs.slides._sldIdLst)), R - 0.6, 7.08, 0.6, 0.25, size=9,
                 color=GRAY, align=PP_ALIGN.RIGHT)
    return s


def _content_bottom(slide):
    """Tepi bawah konten terendah di badan slide (judul dan nomor halaman diabaikan)."""
    bottom = BODY_TOP
    for sh in slide.shapes:
        if sh.name == 'Title' or sh.top is None or sh.height is None:
            continue
        t, b = sh.top / 914400, (sh.top + sh.height) / 914400
        if t >= 7.0 or b > 7.0:
            continue
        bottom = max(bottom, b)
    return bottom


def add_keterangan(slide, teks, x, y, w, h=None, judul=None, size=KET_PT, bold=False):
    """Keterangan singkat DI BADAN SLIDE, sebagai kolom kanan atau kiri di samping grafis/chart/gambar.

    Isinya menjelaskan APA yang ditampilkan grafis (data apa, satuan/periode bila belum di judul
    sumbu, cara membacanya: arti warna, garis target, ukuran titik). 1–2 kalimat atau 2–4 butir.
    JANGAN menambah implikasi, rekomendasi, atau 'jadi apa' yang tidak diberikan user.
    Kesimpulan hanya ditulis bila user memberikannya/memintanya (bold=True untuk kalimat itu).

    teks  : kalimat (str) atau daftar butir (list; butir boleh dict {'text','sub','bold'}).
    judul : judul kolom opsional bergaris bawah PRIMARY (mis. nama grafisnya); None = tanpa judul.
    Tanpa kotak berwarna, tanpa pita aksen."""
    h = h or (BODY_BOTTOM - y)
    yy = add_column_head(slide, judul, x, y, w) if judul else y
    if isinstance(teks, (list, tuple)):
        box = add_bullets(slide, list(teks), x, yy, w, y + h - yy, size=size)
    else:
        box = add_text(slide, teks, x, yy, w, y + h - yy, size=size, color=INK, bold=bold,
                       space_after=6, line=1.1)
    box.name = 'Pesan'
    return box


def add_keterangan_bawah(slide, teks, y=None, x=None, w=None, size=KET_PT, bold=False):
    """Keterangan singkat DI BAWAH KONTEN, selebar bidang isi: garis tipis abu di atas lalu satu-dua
    kalimat rata kiri. Isinya menjelaskan apa yang ditampilkan grafis (lihat add_keterangan);
    bukan implikasi tambahan. Panggil SETELAH konten dibuat. y default: 0,3 in di bawah tepi bawah
    konten terendah (maksimal BODY_BOTTOM - 0,55). y='dasar' untuk selalu di posisi bawah yang sama."""
    x = L if x is None else x
    w = (R - x) if w is None else w
    if y is None:
        y = _content_bottom(slide) + 0.3
        y = min(max(y, BODY_TOP + 0.5), BODY_BOTTOM - 0.55)
    elif y == 'dasar':
        y = BODY_BOTTOM - 0.55
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = OUTLINE
    ln.line.width = Pt(0.75)
    box = add_text(slide, teks, x, y + 0.1, w, BODY_BOTTOM - y - 0.1, size=size, color=INK, bold=bold,
                   line=1.05)
    box.name = 'Pesan'
    return box


# Nama lama tetap bisa dipakai
add_pesan = add_keterangan
add_pesan_bawah = add_keterangan_bawah


def _title_lines(teks, w_in, pt, em=0.53):
    """Perkiraan jumlah baris judul Helvetica Bold (lebar rata-rata karakter ~0,53 em)."""
    import math
    cpl = max(int(w_in * 72 / (pt * em)), 1)
    words, lines, cur = teks.split(), 1, 0
    for wd in words:
        add = len(wd) + (1 if cur else 0)
        if cur + add > cpl:
            lines += 1; cur = len(wd)
        else:
            cur += add
    return lines


FOTO_L = 4.95   # tepi kiri isi pada desain 10/11 (foto di kiri)


def add_source(slide, teks):
    """Tidak dipakai: deck PLN IP tidak memakai baris 'Sumber: ...' (materi internal)."""
    raise ValueError('Baris sumber tidak dipakai. Hapus add_source(); bila asumsi penting, '
                     'tulis sebagai bagian isi (mis. butir di add_keterangan).')


def add_column_head(slide, teks, x, y, w, unit=None, size=14):
    """Judul kolom / judul sumbu: bold + garis bawah PRIMARY. Mengembalikan y isi berikutnya."""
    tw = w * 0.62 if unit else w
    add_text(slide, teks, x, y, tw, HEAD_H, size=size, color=INK, bold=True, anchor=MSO_ANCHOR.BOTTOM)
    if unit:
        add_text(slide, unit, x + tw, y, w - tw, HEAD_H, size=12, color=GRAY,
                 align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.BOTTOM)
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y + HEAD_H + 0.04),
                                    Inches(x + w), Inches(y + HEAD_H + 0.04))
    ln.line.color.rgb = PRIMARY
    ln.line.width = Pt(1.5)
    return y + HEAD_H + 0.16


def add_bullets(slide, items, x, y, w, h, size=13, color=INK):
    """items: 'teks' atau {'text': ..., 'sub': [...], 'bold': True}"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for it in items:
        d = it if isinstance(it, dict) else {'text': it}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(6)
        _bullet(p, '•', 0)
        _fill_runs(p, d['text'], size, color, bold=d.get('bold', False))
        for sub in d.get('sub', []):
            ps = tf.add_paragraph(); ps.space_after = Pt(3)
            _bullet(ps, '–', 1)
            _fill_runs(ps, sub, size - 1, color)
    return box


def _bullet(p, char, level):
    """Bullet asli PowerPoint (bukan karakter yang diketik di dalam teks)."""
    p.level = level
    pPr = p._pPr if p._pPr is not None else p._p.get_or_add_pPr()
    pPr.set('marL', str(int((0.25 + 0.25 * level) * 914400)))
    pPr.set('indent', str(int(-0.2 * 914400)))
    for tag in ('a:buNone', 'a:buChar', 'a:buAutoNum'):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    font = pPr.makeelement(qn('a:buFont'), {'typeface': 'Arial'})
    bu = pPr.makeelement(qn('a:buChar'), {'char': char})
    pPr.append(font); pPr.append(bu)


# ---------- Tabel ----------
def _set_border(cell, edge, color, pt):
    tcPr = cell._tc.get_or_add_tcPr()
    tag = {'top': 'a:lnT', 'bottom': 'a:lnB', 'left': 'a:lnL', 'right': 'a:lnR'}[edge]
    for el in tcPr.findall(qn(tag)):
        tcPr.remove(el)
    ln = tcPr.makeelement(qn(tag), {'w': str(int(pt * 12700)), 'cap': 'flat',
                                    'cmpd': 'sng', 'algn': 'ctr'})
    if color is None:
        ln.append(ln.makeelement(qn('a:noFill'), {}))
    else:
        fill = ln.makeelement(qn('a:solidFill'), {})
        clr = fill.makeelement(qn('a:srgbClr'), {'val': '%02X%02X%02X' % (color[0], color[1], color[2])})
        fill.append(clr); ln.append(fill)
    tcPr.append(ln)


def _cell_fill(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ('a:solidFill', 'a:noFill'):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
    if color is None:
        tcPr.append(tcPr.makeelement(qn('a:noFill'), {}))
    else:
        fill = tcPr.makeelement(qn('a:solidFill'), {})
        fill.append(fill.makeelement(qn('a:srgbClr'),
                    {'val': '%02X%02X%02X' % (color[0], color[1], color[2])}))
        tcPr.append(fill)


def add_axis_table(slide, header, rows, x, y, w, col_w=None, row_h=0.5, size=12,
                   row_axis=True, row_icons=None, icon_d=0.42):
    """Tabel bersumbu: header bold bergaris bawah PRIMARY, antarbaris garis tipis,
    tanpa zebra. Sel boleh string atau dict {'text','bold','na','highlight'}.

    row_icons : daftar nama ikon, satu per baris (None untuk baris tanpa ikon). Ikon putih
                dalam lingkaran PRIMARY diletakkan di kolom sumbu, di kiri label baris
                (pola resmi PLN IP, lihat aturan-slide.md §5.15). Pakai hanya untuk tabel
                kualitatif 3–7 baris dengan row_h tetap; baris yang tingginya ikut
                membesar karena teks panjang akan membuat ikon tidak sejajar."""
    n_rows = len(rows) + (1 if header else 0)
    n_cols = len(header) if header else len(rows[0])
    shape = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(w),
                                   Inches(HEAD_H + row_h * len(rows)))
    tbl = shape.table
    tbl.first_row = bool(header)
    tbl.horz_banding = False
    if col_w:
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Inches(cw)
    r0 = 0
    if header:
        tbl.rows[0].height = Inches(HEAD_H)
        for j, teks in enumerate(header):
            c = tbl.cell(0, j)
            _fmt_cell(c, teks, size + 2, INK, bold=True, anchor=MSO_ANCHOR.BOTTOM)
            _cell_fill(c, None)
            _set_border(c, 'bottom', PRIMARY, 1.5)
            for e in ('top', 'left', 'right'):
                _set_border(c, e, None, 0)
        r0 = 1
    for i, row in enumerate(rows):
        tbl.rows[i + r0].height = Inches(row_h)
        last = i == len(rows) - 1
        for j, cell in enumerate(row):
            d = cell if isinstance(cell, dict) else {'text': cell}
            c = tbl.cell(i + r0, j)
            bold = d.get('bold', False) or (row_axis and j == 0)
            sz = size + 2 if (row_axis and j == 0) else size
            col = GRAY if d.get('na') else INK
            _fmt_cell(c, d['text'], sz, col, bold=bold)
            _cell_fill(c, YELLOW if d.get('highlight') else (MIST if d.get('na') else None))
            _set_border(c, 'bottom', None if last else OUTLINE, 0.75)
            for e in ('top', 'left', 'right'):
                _set_border(c, e, None, 0)
            if row_icons and j == 0 and i < len(row_icons) and row_icons[i]:
                c.margin_left = Inches(icon_d + 0.18)
    if row_icons:
        top = y + (HEAD_H if header else 0)
        for i, nama in enumerate(row_icons[:len(rows)]):
            if nama:
                add_icon_badge(slide, nama, x + 0.06, top + i * row_h + (row_h - icon_d) / 2,
                               d=icon_d)
    return tbl


def _fmt_cell(cell, teks, size, color, bold=False, anchor=MSO_ANCHOR.MIDDLE):
    cell.vertical_anchor = anchor
    cell.margin_left = Inches(0.06); cell.margin_right = Inches(0.06)
    cell.margin_top = Inches(0.03); cell.margin_bottom = Inches(0.03)
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    _fill_runs(p, str(teks), size, color, bold)


# ---------- Chart ----------
def add_bar_chart(slide, categories, values, x, y, w, h, highlight=None, horizontal=True,
                  number_format='#,##0', colors=None):
    """Chart batang satu seri: yang disorot PRIMARY, sisanya abu. highlight = indeks."""
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
    data = CategoryChartData()
    data.categories = list(categories)
    data.add_series('seri', tuple(values), number_format)
    ctype = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = slide.shapes.add_chart(ctype, Inches(x), Inches(y), Inches(w), Inches(h), data)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = number_format
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(11); dl.font.color.rgb = INK; dl.font.name = FONT
    ser = plot.series[0]
    for i, pt in enumerate(ser.points):
        if colors:
            c = colors[i % len(colors)]
        else:
            c = PRIMARY if (highlight is None or i in _as_set(highlight)) else OUTLINE
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = c
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    ca = ch.category_axis
    ca.has_major_gridlines = False
    ca.tick_labels.font.size = Pt(11)
    ca.tick_labels.font.color.rgb = INK
    ca.tick_labels.font.name = FONT
    if horizontal:   # kategori pertama tampil di atas, seperti kebiasaan deck konsultan
        scaling = ca._element.find(qn('c:scaling'))
        if scaling is not None:
            for el in scaling.findall(qn('c:orientation')):
                scaling.remove(el)
            scaling.insert(0, scaling.makeelement(qn('c:orientation'), {'val': 'maxMin'}))
    return ch


def _as_set(v):
    return set(v) if isinstance(v, (list, tuple, set)) else {v}


# ---------- Ikon, lingkaran ikon, ilustrasi ----------
# Skill hanya menyimpan SVG; PNG dirender saat dipakai lewat svg_render.py (cache /tmp/plnip-svg).
import sys as _sys
_sys.path.insert(0, HERE)
from svg_render import svg_to_png, resolve_illustration, cari as _cari  # noqa: E402


def _rgb_hex(c):
    return str(c) if isinstance(c, RGBColor) else c


def icon_path(nama, varian='primary', px=256):
    """varian: 'primary' | 'dark' | 'white' | 'tint' | hex apa pun ('#FFD966' atau RGBColor)."""
    p = os.path.join(HERE, 'icons', 'svg', nama + '.svg')
    if not os.path.exists(p):
        raise FileNotFoundError('Ikon tidak ada: %s (lihat assets/icons/KATALOG.md)' % nama)
    return svg_to_png(p, width=px, color=_rgb_hex(varian))


def illustration_path(nama, accent=PRIMARY, recolor=None, px=1400):
    """accent: warna aksen (currentColor unDraw). recolor: {'#lama': '#baru'} untuk warna lain.
    Palet unDraw (abu gelap/muda) otomatis dipetakan ke DARK/LIGHT PLN IP."""
    p = resolve_illustration(nama)   # folder svg/ (11 flat) lalu undraw-pustaka.json
    m = {k: _rgb_hex(v) for k, v in (recolor or {}).items()}
    return svg_to_png(p, width=px, color=_rgb_hex(accent), mapping=m, undraw=True)


def add_icon(slide, nama, x, y, size=0.34, varian='primary'):
    pic = slide.shapes.add_picture(icon_path(nama, varian), Inches(x), Inches(y),
                                   Inches(size), Inches(size))
    pic.name = 'Ikon ' + nama          # dibaca check_deck.py untuk rekap aset visual
    return pic


def add_icon_badge(slide, nama, x, y, d=0.6, fill=PRIMARY):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.fill.background()
    sh.shadow.inherit = False
    sh.name = 'Lingkaran ikon ' + nama
    s = d * 0.55
    pic = slide.shapes.add_picture(icon_path(nama, 'white'), Inches(x + (d - s) / 2),
                                   Inches(y + (d - s) / 2), Inches(s), Inches(s))
    pic.name = 'Ikon ' + nama
    return sh


def add_icon_list(slide, items, x, y, w, row_h=0.95, d=0.6, size=13):
    """items: {'icon','text','sub'} — 3 sampai 5 butir setara."""
    for i, it in enumerate(items):
        yy = y + i * row_h
        add_icon_badge(slide, it['icon'], x, yy, d)
        add_text(slide, it['text'], x + d + 0.25, yy, w - d - 0.25, 0.4, size=size, bold=True)
        if it.get('sub'):
            add_text(slide, it['sub'], x + d + 0.25, yy + 0.32, w - d - 0.25, row_h - 0.4,
                     size=size - 1, color=GRAY)


def add_illustration(slide, nama, x, y, w=None, h=None, accent=PRIMARY, recolor=None):
    """Beri w ATAU h; rasio asli dijaga. Aturan pakai: aturan-slide.md §5.16."""
    p = illustration_path(nama, accent, recolor)
    if w and h:
        h = None
    pic = slide.shapes.add_picture(p, Inches(x), Inches(y),
                                   Inches(w) if w else None, Inches(h) if h else None)
    pic.name = 'Ilustrasi ' + nama     # dibaca check_deck.py untuk rekap aset visual
    return pic


def fit_illustration(slide, nama, x, y, w, h, accent=PRIMARY, recolor=None, align='center'):
    """Ilustrasi dimuatkan ke dalam kotak (x, y, w, h) tanpa mengubah rasio, lalu
    ditengahkan (align='center') atau dirapatkan ke bawah (align='bottom').
    Cara paling aman mengisi ruang kosong di kolom kanan."""
    from PIL import Image
    p = illustration_path(nama, accent, recolor)
    pw, ph = Image.open(p).size
    scale = min(w / pw, h / ph)
    ww, hh = pw * scale, ph * scale
    xx = x + (w - ww) / 2
    yy = y + (h - hh) / 2 if align == 'center' else y + h - hh
    pic = slide.shapes.add_picture(p, Inches(xx), Inches(yy), Inches(ww), Inches(hh))
    pic.name = 'Ilustrasi ' + nama
    return pic


def add_section_slide(prs, bab, aktif=None, ilustrasi=None, topik='Daftar Isi', subtopik='',
                      accent=PRIMARY):
    """Daftar isi (aktif=None) atau pembatas bab (aktif = indeks bab, mulai 1) bergaya
    content tracker: daftar bab satu kolom di kiri, satu ilustrasi di kanan.

    bab       : daftar judul bab, 3–7 butir, sama persis dengan topik judul slide di bab itu
    aktif     : nomor bab yang sedang dibuka; bab lain tampil abu, bab aktif DARK tebal
                dengan pita TINT. Judul slide otomatis menjadi 'Bab n' + nama bab.
    ilustrasi : nama di assets/illustrations/KATALOG.md (lihat tabel 'Ilustrasi per jenis slide')
    """
    if aktif:
        s = add_slide(prs, 'Bab %d' % aktif, subtopik or bab[aktif - 1])
    else:
        s = add_slide(prs, topik, subtopik)
    n = len(bab)
    row_h = min(0.78, (BODY_BOTTOM - BODY_TOP - 0.3) / max(n, 1))
    y0 = BODY_TOP + (BODY_BOTTOM - BODY_TOP - row_h * n) / 2
    w = LEFT_W
    for i, teks in enumerate(bab, 1):
        y = y0 + (i - 1) * row_h
        on = aktif is None or aktif == i
        if aktif == i:
            band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(L - 0.1), Inches(y + 0.04),
                                      Inches(w + 0.1), Inches(row_h - 0.08))
            band.fill.solid(); band.fill.fore_color.rgb = TINT; band.line.fill.background()
            band.shadow.inherit = False
            band.name = 'Sorotan bab aktif'
        d = 0.44
        c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(L + 0.05), Inches(y + (row_h - d) / 2),
                               Inches(d), Inches(d))
        c.fill.solid(); c.fill.fore_color.rgb = PRIMARY if on else OUTLINE
        c.line.fill.background(); c.shadow.inherit = False
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p0 = tf.paragraphs[0]; p0.alignment = PP_ALIGN.CENTER
        _run(p0, str(i), 13, WHITE, bold=True)
        add_text(s, teks, L + 0.7, y, w - 0.8, row_h, size=18,
                 color=(DARK if on else GRAY), bold=(aktif == i or aktif is None),
                 anchor=MSO_ANCHOR.MIDDLE)
    if ilustrasi:
        fit_illustration(s, ilustrasi, RIGHT_X + 0.5, BODY_TOP + 0.2, 4.6,
                         BODY_BOTTOM - BODY_TOP - 0.4, accent=accent)
    return s


# Kata kunci -> kandidat ilustrasi. Dipakai lewat saran_ilustrasi('hukum') dsb.
# Isi harus sama dengan tabel di assets/illustrations/KATALOG.md.
TEMA_ILUSTRASI = {
    'rapat':      ['brainstorming', 'conference-call', 'pitching', 'meeting', 'business-decisions', 'presentation', 'online-meeting', 'business-chat'],
    'keputusan':  ['informed-decision', 'right-direction', 'prioritise', 'business-decisions', 'contract-signed', 'agreement', 'target', 'team-goals'],
    'tujuan':     ['target', 'goal', 'team-goals', 'visionary-technology'],
    'latar':      ['researching', 'mind-map', 'map', 'location-search'],
    'lahan':      ['destinations', 'map', 'location-search', 'route-planning', 'property-agreement', 'environment'],
    'lokasi':     ['map', 'location-search', 'route-planning', 'kota-jaringan'],
    'hukum':      ['legal-counsel', 'digital-signature', 'contract', 'signed-document', 'document-warning', 'property-agreement', 'agreement'],
    'kontrak':    ['handshake-deal', 'contract-signed', 'contract', 'agreement', 'business-deal', 'signed-document'],
    'dokumen':    ['documents', 'document-ready', 'document-analysis', 'add-document', 'report'],
    'risiko':     ['warning', 'inspection', 'security', 'document-warning', 'checklist', 'maintenance'],
    'prasyarat':  ['pending-approval', 'approve', 'verified', 'checklist', 'document-ready', 'contract-signed'],
    'analisis':   ['data-analysis', 'analytics', 'business-analytics', 'visual-data', 'statistic-chart'],
    'keuangan':   ['budgeting', 'financial-data', 'investor-update', 'finance', 'investment', 'investment-data', 'revenue-analysis', 'personal-finance'],
    'biaya':      ['finance', 'revenue-analysis', 'investment-data', 'pie-chart'],
    'kinerja':    ['performance-overview', 'performance-comparison', 'dashboard', 'growth-chart'],
    'perbandingan': ['performance-comparison', 'business-analytics', 'heatmap'],
    'teknis':     ['inspection', 'engineering-team', 'construction-workers', 'maintenance', 'factory'],
    'konstruksi': ['construction-workers', 'engineering-team', 'project-team'],
    'proyek':     ['project-team', 'organizing-projects', 'project-completed', 'team-assignment'],
    'tahapan':    ['alur-tahapan', 'creation-process', 'route-planning', 'organizing-projects'],
    'tim':        ['team', 'teamwork', 'team-collaboration', 'team-effort', 'engineering-team'],
    'lingkungan': ['environment', 'environmental-study', 'wind-turbines'],
    'digital':    ['dashboard', 'real-time-analytics', 'artificial-intelligence', 'app-data'],
    'plts':       ['plts', 'visionary-technology'],
    'pltb':       ['pltb', 'wind-turbine', 'wind-turbines'],
    'pltu':       ['pltu', 'factory'],
    'plta':       ['plta'],
    'gas':        ['pipa-gas', 'factory'],
    'bess':       ['bess', 'electricity'],
    'jaringan':   ['transmisi', 'kota-jaringan', 'electricity'],
    'listrik':    ['electricity', 'transmisi', 'kota-jaringan'],
    'kendaraan listrik': ['electric-car'],
    'strategi':   ['master-plan', 'five-year-plan', 'idea-to-plan', 'business-plan', 'social-strategy'],
    'solusi':     ['our-solution', 'solution-mindset', 'problem-solving', 'ideas'],
    'penutup':    ['team-effort', 'project-completed', 'business-deal'],
}


def saran_ilustrasi(kata):
    """Kandidat pilihan kurasi untuk kata kunci Indonesia ('hukum', 'lahan', 'keuangan', ...)."""
    kata = kata.lower()
    hasil = []
    for k, v in TEMA_ILUSTRASI.items():
        if k in kata or kata in k:
            hasil += [n for n in v if n not in hasil]
    return hasil


def cari_ilustrasi(kata, n=15):
    """Cari di seluruh pustaka unDraw (768 ilustrasi) dengan kata kunci Inggris
    ('contract', 'solar', 'meeting team'). Pakai bila saran_ilustrasi() kurang pas."""
    return _cari(kata, n)


def add_causal_triangle(slide, x_center, y_center, size=0.24, color=PRIMARY):
    sh = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                Inches(x_center - size / 2), Inches(y_center - size / 2),
                                Inches(size), Inches(size))
    sh.rotation = 90
    sh.fill.solid(); sh.fill.fore_color.rgb = color; sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def add_chevrons(slide, labels, x, y, w, h=0.55, muted=(), size=12):
    overlap = 0.12
    n = len(labels)
    sw = (w + overlap * (n - 1)) / n
    for i, lab in enumerate(labels):
        shape = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        sh = slide.shapes.add_shape(shape, Inches(x + i * (sw - overlap)), Inches(y),
                                    Inches(sw), Inches(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = OUTLINE if i in muted else PRIMARY
        sh.line.fill.background()
        sh.shadow.inherit = False
        tf = sh.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        _fill_runs(p, lab, size, WHITE, bold=True)
    return sw - overlap


# ---------- Model grafis tambahan ----------
from plnip_grafis import *  # noqa: E402,F401,F403
